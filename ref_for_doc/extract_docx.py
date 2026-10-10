import sys
import subprocess
import os

def install_docx():
    try:
        import docx
    except ImportError:
        subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])

install_docx()
import docx

try:
    doc = docx.Document(r'f:\NIDM-SRM Drone Project Docs\ref_for_doc\1988-2023.docx')
    with open(r'f:\NIDM-SRM Drone Project Docs\ref_for_doc\docx_dump.txt', 'w', encoding='utf-8') as f:
        for p in doc.paragraphs:
            if p.text.strip():
                f.write(p.text.strip() + '\n')
        # Extract tables as well
        for table in doc.tables:
            for row in table.rows:
                f.write(' | '.join([cell.text.strip() for cell in row.cells]) + '\n')
            f.write('\n')
except Exception as e:
    with open(r'f:\NIDM-SRM Drone Project Docs\ref_for_doc\docx_dump.txt', 'w', encoding='utf-8') as f:
        f.write('Error: ' + str(e))
